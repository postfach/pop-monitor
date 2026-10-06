"""Material Instructions workbooks laid out like the real ones (for tests and
the mock Graph). Two layouts seen in SharePoint: the standard template (Pals,
dates in the creative column headers), and the variant with unnamed size
columns, a row of group labels and day-parting notes (Anchor)."""
import io
from datetime import datetime

import openpyxl

SPECS = [("LUMO-Beach", "576 x 396"), ("LUMO-Curletts In", "432 x 864"), ("LUMO-Curletts Out", "432 x 864"),
         ("LUMO-Albany Expressway In", "1188 x 288"), ("LUMO-Albany Expressway Out", "1188 x 288"),
         ("LUMO-Khyber", "1188 x 288"), ("LUMO-Customhouse Quay", "576 x 396"), ("LUMO-Tāmaki", "1368 x 324")]


def _top(ws, cid, name, start, end, updated):
    for i, (k, v) in enumerate([("Contract #", cid), ("Contract Name", name), ("Contract Start", start),
                                ("Contract End", end), ("Instructions last updated", updated)], start=2):
        ws.cell(i, 2, k)
        ws.cell(i, 3, v)


def _specs(wb):
    ws = wb.create_sheet("2. Screen Specs")
    ws.append(["Screen ID", "Screen Spec", "Ratio", "Screen Physical Size ", "Ad Length"])
    for sid, spec in SPECS:
        ws.append([sid, spec, "", "", 8])


def standard(cid, name, rows, headers=("Creative #1 File Name", "Creative #2 File Name"), updated="17.08.2026"):
    """rows: (site, start, end, type, [file names...], w, h)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1. Material Instructions "
    _top(ws, cid, name, rows[0][1], rows[-1][2], updated)
    ws.cell(9, 6, "File name must match supplied creative")
    hdr = ["Site Name", "Booking Start", "Booking End", "Booking Type"]
    for h in headers:
        hdr += [h, "Rotation %"]
    hdr += ["Pixel Width", "Pixel Height", "Creative Start Date", "Creative End Date", "Start Time", "End Time", "Notes"]
    for j, v in enumerate(hdr, start=2):
        ws.cell(10, j, v)
    r = 11
    for site, start, end, typ, files, w, h in rows:
        vals = [site, start, end, typ]
        for k in range(len(headers)):
            vals += [files[k] if k < len(files) else None, None]
        vals += [w, h, None, None, None, None, None]
        for j, v in enumerate(vals, start=2):
            ws.cell(r, j, v)
        r += 1
    ws.cell(r + 1, 2, "Key info")
    ws.cell(r + 2, 2, "Refer to tab  2. Screen Specs  for screen specific dimensions. ")
    _specs(wb)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def anchor_style(cid, name, rows, updated="15.09.2026"):
    """rows: (site, start, end, type, w, h, [(file, rotation)...], notes) with real date cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1. Material Instructions "
    _top(ws, cid, name, rows[0][1], rows[-1][2], updated)
    hdr = ["Site Name", "Booking Start", "Booking End", "Booking Type", None, None]
    for k in range(1, 7):
        hdr += [f"Creative #{k} File Name", "Rotation %"]
    hdr += ["Pixel Width", "Pixel Height", "Creative Start Date", "Creative End Date", "Start Time", "End Time", "Notes"]
    for j, v in enumerate(hdr, start=2):
        ws.cell(9, j, v)
    for j, v in ((8, "REAL GOOD"), (12, "REAL BENEFIT"), (16, "WAKE UP TO MILK")):
        ws.cell(10, j, v)
    r = 11
    for site, start, end, typ, w, h, files, notes in rows:
        s, e = datetime.strptime(start, "%d.%m.%Y"), datetime.strptime(end, "%d.%m.%Y")
        vals = [site, s, e, typ, w, h]
        for k in range(6):
            vals += list(files[k]) if k < len(files) else [None, None]
        vals += [w, h, s, e, None, None, notes]
        for j, v in enumerate(vals, start=2):
            ws.cell(r, j, v)
        r += 1
    _specs(wb)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
