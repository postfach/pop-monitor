"""Reading LUMO Material Instructions workbooks and matching file names.

Kept separate from the service so it can be unit tested without Graph.

A Material Instructions (MI) workbook is generated from a template, with a
"1. Material Instructions" sheet:

    Site Name | Booking Start | Booking End | Booking Type | [W] [H] |
    Creative #1 File Name[: 5th Oct - 7th October] | Rotation % | Creative #2 ... |
    Pixel Width | Pixel Height | Creative Start Date | Creative End Date |
    Start Time | End Time | Notes

and a "2. Screen Specs" sheet (Screen ID, Screen Spec "1188 x 288", ...). Real
files vary: extra unnamed columns, a row of group labels under the header,
rotation as "15%", "50" or 0.15, dates as cells or "05.10.2026" text, file
names with or without the extension and with typos, and some campaigns use a
"Booking Update" sheet laid out the same way. Everything here is tolerant of
that and never raises on a bad cell.
"""
import difflib
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
IMAGE_EXT = (".jpg", ".jpeg", ".png")
VIDEO_EXT = (".mp4", ".mov", ".m4v", ".webm", ".avi", ".mpg", ".mpeg")

_WORDS = {"road": "rd", "street": "st", "saint": "st", "mount": "mt", "avenue": "ave"}


def fold(s):
    """Lower case, accents removed (Tāmaki -> tamaki), punctuation to spaces."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def screen_norm(name):
    """Key for a screen name as written in Vault, an MI ("LUMO-Blenheim Rd") or a file name."""
    s = fold(name)
    s = re.sub(r"^lumo\s+", "", s)
    return " ".join(_WORDS.get(w, w) for w in s.split())


def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", fold(name)).strip("-") or "screen"


def stem(name):
    n = str(name or "").strip()
    low = n.lower()
    for ext in IMAGE_EXT + VIDEO_EXT:
        if low.endswith(ext):
            return n[: -len(ext)]
    return n


def name_norm(name):
    """Comparable form of a file name: no extension, lower case, only letters and digits."""
    return re.sub(r"[^a-z0-9]", "", fold(stem(name)))


def size_in_name(name):
    m = re.search(r"(?<!\d)(\d{3,4})\s*[x×X]\s*(\d{3,4})(?!\d)", str(name or ""))
    return (int(m.group(1)), int(m.group(2))) if m else None


def same_size(actual, want):
    """Artwork fits a screen: the exact size, or the same shape at up to 3x
    (artwork is often supplied at 2x, or a pixel over: 577x397 for 576x396)."""
    if not actual or not want:
        return None
    (w, h), (W, H) = actual, want
    if (w, h) == (W, H):
        return True
    if not (w and h and W and H):
        return False
    return abs((w / h) / (W / H) - 1) <= 0.012 and 0.97 <= w / W <= 3.05


def kind_of(name):
    low = str(name or "").lower()
    if low.endswith(IMAGE_EXT):
        return "image"
    if low.endswith(VIDEO_EXT):
        return "video"
    return None


# ---------------------------------------------------------------- dates

def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})(?:\s.*)?", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y += 2000 if y < 100 else 0
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    m = re.fullmatch(r"(\d{4})-(\d\d)-(\d\d).*", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _day_month(txt, year):
    """'5th Oct', 'October 8', '8/10' -> date in `year`, or None."""
    t = fold(txt)
    m = re.search(r"(\d{1,2})\s*(?:st|nd|rd|th)?\s+([a-z]{3,})", t)
    if m and m.group(2)[:3] in MONTHS:
        d, mo = int(m.group(1)), MONTHS[m.group(2)[:3]]
    else:
        m = re.search(r"([a-z]{3,})\s+(\d{1,2})", t)
        if m and m.group(1)[:3] in MONTHS:
            d, mo = int(m.group(2)), MONTHS[m.group(1)[:3]]
        else:
            m = re.search(r"(\d{1,2})\s+(\d{1,2})(?:\s+(\d{2,4}))?", t)
            if not m:
                return None
            d, mo = int(m.group(1)), int(m.group(2))
            if m.group(3):
                year = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
    try:
        return date(year, mo, d)
    except ValueError:
        return None


def header_range(header, year):
    """Dates in a creative column header: 'Creative #1 File Name: 5th Oct - 7th October'."""
    if not header or not year:
        return None
    h = str(header)
    parts = re.split(r":", h, maxsplit=1)
    tail = parts[1] if len(parts) > 1 else re.sub(r"(?i)creative\s*#?\s*\d*\s*(file\s*name)?", "", h)
    bits = re.split(r"\s+(?:-|–|—|to|until)\s+", tail.strip(), maxsplit=1)
    if len(bits) != 2:
        return None
    a, b = _day_month(bits[0], year), _day_month(bits[1], year)
    if not a or not b:
        return None
    if b < a:                        # "28th Dec - 3rd Jan"
        b = b.replace(year=b.year + 1)
    return a, b


def rotation(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(v * 100) if 0 < v <= 1 else round(v)
    m = re.search(r"(\d+(?:\.\d+)?)", str(v))
    if not m:
        return None
    x = float(m.group(1))
    return round(x * 100) if "%" not in str(v) and 0 < x <= 1 else round(x)


def time_text(v):
    if v is None or v == "":
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    s = str(v).strip()
    return s or None


# ---------------------------------------------------------------- workbooks

def _cells(ws, max_rows=2000, max_cols=60):
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, max_rows),
                            max_col=min(ws.max_column or 0, max_cols), values_only=True):
        rows.append(list(row))
    return rows


def _find_header(rows, words):
    for i, row in enumerate(rows[:60]):
        cells = [fold(c) for c in row]
        if any(c in words for c in cells):
            return i, cells
    return None, None


def _col(cells, *tests, start=0):
    for j in range(start, len(cells)):
        c = cells[j]
        if c and any(t(c) for t in tests):
            return j
    return None


def parse_specs(rows):
    """'2. Screen Specs' sheet -> {screen_norm: (w, h, name)}."""
    i, cells = _find_header(rows, {"screen id"})
    if i is None:
        return {}
    cid = cells.index("screen id")
    cspec = _col(cells, lambda c: c == "screen spec", lambda c: "spec" in c and "screen" in c)
    out = {}
    for row in rows[i + 1:]:
        name = row[cid] if cid < len(row) else None
        spec = row[cspec] if cspec is not None and cspec < len(row) else None
        sz = size_in_name(spec)
        if name and sz:
            raw = re.sub(r"(?i)^lumo\s*[-_ ]\s*", "", str(name)).strip()
            out[screen_norm(raw)] = (sz[0], sz[1], raw)
    return out


def parse_rows(rows):
    """Booking rows of an MI-style sheet. Returns a list of dicts or [] if this
    isn't one."""
    i, cells = _find_header(rows, {"site name", "site", "screen", "screen name"})
    if i is None:
        return []
    csite = _col(cells, lambda c: c in ("site name", "site", "screen", "screen name"))
    cstart = _col(cells, lambda c: c in ("booking start", "start date", "start", "booking start date"),
                  lambda c: c.startswith("booking start"))
    cend = _col(cells, lambda c: c in ("booking end", "end date", "end", "booking end date"),
                lambda c: c.startswith("booking end"))
    ctype = _col(cells, lambda c: c.startswith("booking type"))
    cw = _col(cells, lambda c: c.startswith("pixel width"))
    ch = _col(cells, lambda c: c.startswith("pixel height"))
    ccs = _col(cells, lambda c: c.startswith("creative start"))
    cce = _col(cells, lambda c: c.startswith("creative end"))
    cts = _col(cells, lambda c: c == "start time")
    cte = _col(cells, lambda c: c == "end time")
    cnotes = _col(cells, lambda c: c.startswith("notes"))
    creative_cols = []
    for j, c in enumerate(cells):
        if not c or j in (ccs, cce):
            continue
        if re.match(r"creative\s*\d*\s*(file|name|$)", c) or ("creative" in c and "file" in c) or c.startswith("file name"):
            rot = None
            for k in (j + 1, j + 2):
                if k < len(cells) and cells[k] and "rotation" in cells[k]:
                    rot = k
                    break
            creative_cols.append((j, rot, rows[i][j]))
    if csite is None or not creative_cols:
        return []
    out = []
    for row in rows[i + 1:]:
        get = lambda j: row[j] if j is not None and j < len(row) else None
        site = get(csite)
        if not site or not str(site).strip():
            continue
        start, end = to_date(get(cstart)), to_date(get(cend))
        if not start and not end:
            continue                      # group labels, notes, key info
        year = (start or end).year
        w, h = get(cw), get(ch)
        try:
            size = (int(w), int(h)) if w and h else None
        except (TypeError, ValueError):
            size = None
        cstart_d, cend_d = to_date(get(ccs)), to_date(get(cce))
        creatives = []
        for j, rot, header in creative_cols:
            v = get(j)
            if v is None or not str(v).strip():
                continue
            for name in re.split(r"\s*[\n;]\s*", str(v).strip()):
                if not name or fold(name) in ("tbc", "na", "n a", "none", "same as above"):
                    continue
                rng = header_range(header, year)
                creatives.append({"name": name, "rotation": rotation(get(rot)) if rot is not None else None,
                                  "from": (rng[0] if rng else cstart_d or start),
                                  "to": (rng[1] if rng else cend_d or end),
                                  "column": str(header or "").strip()})
        ts, te = time_text(get(cts)), time_text(get(cte))
        out.append({"site": str(site).strip(), "screen": screen_norm(re.sub(r"(?i)^lumo\s*[-_ ]\s*", "", str(site).strip())),
                    "start": start or end, "end": end or start,
                    "type": str(get(ctype) or "").strip(), "size": size, "creatives": creatives,
                    "hours": f"{ts}-{te}" if ts and te else None,
                    "notes": str(get(cnotes) or "").strip()})
    return out


def read_workbook(data):
    """MI workbook bytes -> {"rows": [...], "specs": {...}, "updated": date|None}."""
    import openpyxl                       # imported here so tests of the rest don't need it
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows, specs, updated = [], {}, None
    for ws in wb.worksheets:
        cells = _cells(ws)
        title = fold(ws.title)
        if "spec" in title:
            specs.update(parse_specs(cells))
            continue
        got = parse_rows(cells)
        if got:
            rows.extend(got)
        for r in cells[:12]:
            if r and fold(r[0] if r[0] is not None else (r[1] if len(r) > 1 else "")).startswith("instructions last updated"):
                for v in r[1:]:
                    d = to_date(v)
                    if d:
                        updated = max(updated, d) if updated else d
        if not specs:
            specs.update(parse_specs(cells))
    wb.close()
    return {"rows": rows, "specs": specs, "updated": updated}


# ---------------------------------------------------------------- matching

def match_file(wanted, files, size=None):
    """Best file for a name from the MI. `files` are dicts with "name" (and
    optionally "width"/"height"). Returns (file, how) or (None, None)."""
    want = name_norm(wanted)
    if not want:
        return None, None
    want_ext = kind_of(wanted)

    def ok_size(f):
        s = (f.get("width"), f.get("height")) if f.get("width") else size_in_name(f["name"])
        return not size or not s or bool(same_size(s, size)) or same_size(size_in_name(f["name"]), size)

    exact = [f for f in files if name_norm(f["name"]) == want]
    if exact:
        exact.sort(key=lambda f: (not ok_size(f), (kind_of(f["name"]) != want_ext) if want_ext else False,
                                  f.get("rank", 0)))
        return exact[0], "exact"
    best, score = None, 0.0
    for f in files:
        n = name_norm(f["name"])
        if not n:
            continue
        r = difflib.SequenceMatcher(None, want, n).ratio()
        if want in n or n in want:
            r = max(r, 0.9 if min(len(n), len(want)) >= 12 else r)
        if not ok_size(f):
            r -= 0.15
        if r > score:
            best, score = f, r
    if best is not None and score >= 0.86:
        return best, f"close ({score:.2f})"
    return None, None


def names_screen(filename, screen, others=()):
    """Does a file name name this screen (and not a longer screen name that
    contains it, e.g. 'Curletts' vs 'Curletts In')?"""
    f = " " + screen_norm(stem(filename)) + " "
    key = screen_norm(screen)
    variants = {key, key.replace(" in", "").replace(" out", "")} if key else set()
    hit = any(v and f" {v} " in f for v in variants)
    if not hit:
        # typo-tolerant on the screen's first word ("Curlettes" for "Curletts")
        words = key.split()
        if words and len(words[0]) >= 5:
            for tok in f.split():
                if difflib.SequenceMatcher(None, tok, words[0]).ratio() >= 0.85:
                    rest = " ".join(words[1:])
                    if not rest or f" {rest} " in f.split(tok, 1)[1] + " ":
                        hit = True
                        break
    if not hit:
        return False
    for o in others:                     # a more specific screen name also matches: it's that one
        ok = screen_norm(o)
        if ok != key and key in ok and f" {ok} " in f:
            return False
    # "Curletts In" must not take "Curletts Out" files
    for a, b in ((" in", " out"), (" out", " in")):
        if key.endswith(a) and f" {key[: -len(a)]}{b} " in f:
            return False
    return True


def week_folder_date(path):
    """'.../w.c 05.10.26/...' or 'WC 5 Oct' -> date of that week, else None."""
    for seg in reversed(str(path).split("/")):
        m = re.search(r"(?i)\bw\s*/?\s*\.?\s*c\b\.?\s*(.+)", seg)
        if m:
            d = to_date(re.sub(r"\s+", "", m.group(1))[:10]) or _day_month(m.group(1), date.today().year)
            if d:
                return d
    return None


def monday(d):
    return d - timedelta(days=d.weekday())
